"""Import a master resume from an uploaded .docx or .pdf file.

Text extraction is deterministic (python-docx / pypdf); turning that raw
text into the structured ResumeDocument shape is one LLM call, since
resume layouts vary too much for reliable regex/heuristic parsing. The
result is never saved directly - it's handed back to the Resume page's
existing YAML editor for you to review before Save (which re-validates
against the schema anyway), since this becomes Agent 2b's fabrication-check
ground truth and extraction mistakes matter.
"""

from __future__ import annotations

import io
from pathlib import Path

import yaml
from docx import Document
from pypdf import PdfReader

from jobapply.llm.client import get_client_for
from jobapply.resume.schema import ResumeDocument

_PROMPT_PATH = Path(__file__).parent / "prompts" / "import_system.md"
_MIN_TEXT_LENGTH = 200  # below this, extraction almost certainly failed (e.g. a scanned/image PDF)


class ResumeImportError(Exception):
    """User-facing import failure: bad file type, unreadable content, or extraction too short to be real."""


def _system_prompt() -> str:
    return _PROMPT_PATH.read_text()


def extract_text_from_docx(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    return "\n".join(parts)


def extract_text_from_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_text(filename: str, content: bytes) -> str:
    lowered = filename.lower()
    if lowered.endswith(".docx"):
        return extract_text_from_docx(content)
    if lowered.endswith(".pdf"):
        return extract_text_from_pdf(content)
    raise ResumeImportError(f"Unsupported file type: {filename}. Upload a .docx or .pdf file.")


def structure_resume_text(raw_text: str) -> ResumeDocument:
    client, model = get_client_for("resume_import")
    return client.complete_json(system=_system_prompt(), user=raw_text, response_schema=ResumeDocument, model=model)


def import_resume_yaml(filename: str, content: bytes) -> str:
    """Extracts text, structures it via one LLM call, returns YAML text -
    never writes to disk. Raises ResumeImportError for user-facing
    failures (bad file type, unreadable/too-short extraction)."""
    raw_text = extract_text(filename, content).strip()
    if len(raw_text) < _MIN_TEXT_LENGTH:
        raise ResumeImportError(
            "Couldn't extract enough readable text from this file. If it's a scanned/image PDF "
            "(no selectable text), try a text-based export instead, or paste your resume in manually below."
        )

    resume = structure_resume_text(raw_text)
    return yaml.safe_dump(resume.model_dump(mode="json"), sort_keys=False, allow_unicode=True)
