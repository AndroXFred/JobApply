from __future__ import annotations

import io

import pytest
import yaml
from docx import Document
from fpdf import FPDF

from jobapply.resume.importer import (
    ResumeImportError,
    extract_text,
    extract_text_from_docx,
    extract_text_from_pdf,
    import_resume_yaml,
)
from jobapply.resume.schema import ContactInfo, ExperienceEntry, ResumeDocument

RESUME_TEXT_LINES = [
    "Jane Doe",
    "jane.doe@example.com",
    "Senior Backend Engineer with 8 years of experience building distributed systems.",
    "Experience: Acme Corp, Senior Engineer, 2020-present.",
    "Led backend rewrite reducing latency by 30 percent across the platform.",
    "Skills: Python, FastAPI, PostgreSQL, Kubernetes.",
    "Education: State University, B.S. Computer Science, 2018.",
]


def _make_docx_bytes(paragraphs: list[str]) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_pdf_bytes(lines: list[str]) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in lines:
        pdf.cell(0, 10, text=line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def test_extract_text_from_docx_reads_paragraphs():
    text = extract_text_from_docx(_make_docx_bytes(RESUME_TEXT_LINES))
    assert "Jane Doe" in text
    assert "Kubernetes" in text


def test_extract_text_from_docx_reads_table_cells():
    doc = Document()
    doc.add_paragraph("Jane Doe")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Python, Kubernetes"
    buf = io.BytesIO()
    doc.save(buf)

    text = extract_text_from_docx(buf.getvalue())
    assert "Python, Kubernetes" in text


def test_extract_text_from_pdf_reads_text():
    text = extract_text_from_pdf(_make_pdf_bytes(RESUME_TEXT_LINES))
    assert "Jane Doe" in text
    assert "Kubernetes" in text


def test_extract_text_dispatches_by_extension():
    assert "Jane Doe" in extract_text("resume.docx", _make_docx_bytes(RESUME_TEXT_LINES))
    assert "Jane Doe" in extract_text("resume.pdf", _make_pdf_bytes(RESUME_TEXT_LINES))


def test_extract_text_rejects_unsupported_extension():
    with pytest.raises(ResumeImportError):
        extract_text("resume.txt", b"whatever")


def test_import_resume_yaml_rejects_too_short_extraction():
    with pytest.raises(ResumeImportError):
        import_resume_yaml("resume.docx", _make_docx_bytes(["hi"]))


def test_import_resume_yaml_rejects_unsupported_type():
    with pytest.raises(ResumeImportError):
        import_resume_yaml("resume.txt", b"whatever content, long enough to pass a length check if there were one")


class _FakeClient:
    def __init__(self, resume: ResumeDocument):
        self._resume = resume

    def complete_json(self, *, system, user, response_schema, model):
        return self._resume


def test_import_resume_yaml_produces_valid_yaml(temp_db, monkeypatch):
    fake_resume = ResumeDocument(
        contact=ContactInfo(first_name="Jane", last_name="Doe", email="jane.doe@example.com"),
        summary="Senior Backend Engineer with 8 years of experience.",
        skills=["Python", "FastAPI"],
        experience=[
            ExperienceEntry(
                company="Acme Corp", title="Senior Engineer", start_date="2020", end_date="present",
                bullets=["Led backend rewrite."],
            )
        ],
    )
    monkeypatch.setattr(
        "jobapply.resume.importer.get_client_for", lambda agent: (_FakeClient(fake_resume), "fake-model")
    )

    yaml_text = import_resume_yaml("resume.docx", _make_docx_bytes(RESUME_TEXT_LINES))

    # round-trips through the real schema, same validation a manual paste would hit
    parsed = ResumeDocument.model_validate(yaml.safe_load(yaml_text))
    assert parsed.contact.first_name == "Jane"
    assert "Python" in parsed.skills
